"""Generira testnu dokumentaciju izmišljene Banke Demo d.d.

  data/raw/pdf/*.pdf    - 2 PDF dokumenta
  data/raw/docx/*.docx  - 2 Word dokumenta

Web stranica banke gradi se posebno: scripts/build_demo_site.py

Pokreće se jednom; generirani dokumenti su već u repozitoriju.
Brojke su usklađene s data/rates.json.

    python -m scripts.generate_demo_docs
"""

from pathlib import Path

from docx import Document
from fpdf import FPDF

from bankbot.config import RAW_DIR

# Dokument = naslov + lista blokova: ("h2", tekst), ("p", tekst), ("ul", [stavke]), ("table", [redovi])

GOTOVINSKI = (
    "Opći uvjeti za gotovinske nenamjenske kredite",
    [
        ("p", "Banka Demo d.d. (u daljnjem tekstu: Banka) odobrava gotovinske nenamjenske kredite fizičkim osobama. "
              "Ovaj dokument opisuje uvjete kreditiranja koji vrijede od 1. rujna 2026."),
        ("h2", "Osnovne karakteristike kredita"),
        ("ul", [
            "Valuta kredita: isključivo EUR. Banka ne odobrava kredite u drugim valutama niti s valutnom klauzulom.",
            "Iznos kredita: od 1.000,00 EUR do 40.000,00 EUR.",
            "Rok otplate: od 12 do 120 mjeseci (od 1 do 10 godina).",
            "Kamatna stopa: fiksna za cijelo vrijeme otplate kredita.",
            "Otplata: u jednakim mjesečnim anuitetima.",
            "Namjena: nenamjenski, sredstva se isplaćuju na tekući račun klijenta.",
        ]),
        ("h2", "Tko može podnijeti zahtjev"),
        ("p", "Gotovinski kredit mogu ugovoriti državljani Republike Hrvatske i osobe koje nisu državljani Republike Hrvatske, "
              "a imaju odobren stalni ili privremeni boravak u Republici Hrvatskoj i stalna mjesečna primanja u Republici Hrvatskoj. "
              "Za osobe koje nisu državljani Republike Hrvatske kamatna stopa je viša."),
        ("ul", [
            "Klijent mora imati najmanje 18 godina.",
            "Kredit mora biti otplaćen do navršene 75. godine života klijenta.",
            "Klijent mora imati stalno zaposlenje najmanje 6 mjeseci ili primati mirovinu.",
            "Nije potrebno primati plaću u Banci Demo; kamatna stopa ne ovisi o tome u kojoj banci klijent prima plaću.",
        ]),
        ("h2", "Kamatne stope"),
        ("p", "Kamatna stopa ovisi o roku otplate, iznosu kredita i o tome je li klijent državljanin Republike Hrvatske. "
              "Za državljane RH fiksne kamatne stope kreću se od 5,99 % do 7,19 % godišnje, a za osobe koje nisu državljani RH "
              "od 7,19 % do 8,39 % godišnje. Točnu stopu i mjesečnu ratu za željeni iznos i rok možete provjeriti u kalkulatoru kredita."),
        ("h2", "Potrebna dokumentacija"),
        ("ul", [
            "Zahtjev za kredit (ispunjava se u poslovnici ili u mobilnom bankarstvu).",
            "Važeća osobna iskaznica ili putovnica; za strane državljane i dozvola boravka.",
            "Potvrda poslodavca o zaposlenju i visini primanja, ne starija od 30 dana.",
            "Izvadak s tekućeg računa za posljednja 3 mjeseca ako primanja nisu na računu u Banci Demo.",
            "Za umirovljenike: posljednji odrezak mirovine.",
        ]),
        ("h2", "Naknade"),
        ("ul", [
            "Naknada za obradu kredita: 1 % iznosa kredita, najmanje 50,00 EUR, najviše 300,00 EUR.",
            "Naknada za prijevremenu otplatu: 1 % prijevremeno otplaćenog iznosa ako je preostali rok otplate dulji od godine dana, "
            "odnosno 0,5 % ako je kraći. Za iznose do 2.000,00 EUR godišnje prijevremena otplata je bez naknade.",
            "Izrada otplatnog plana na zahtjev: bez naknade.",
        ]),
        ("h2", "Osiguranje kredita"),
        ("p", "Za kredite do 15.000,00 EUR dovoljna je zadužnica klijenta. Za kredite iznad 15.000,00 EUR potrebna je zadužnica "
              "i jedan jamac platac ili polica osiguranja otplate kredita. Osiguranje otplate kredita nije obvezno, ali se preporučuje."),
        ("h2", "Isplata i otplata"),
        ("p", "Kredit se isplaćuje na tekući račun klijenta najkasnije dva radna dana nakon potpisa ugovora. "
              "Prva rata dospijeva mjesec dana nakon isplate. Klijent može odabrati datum dospijeća rate: 1., 10. ili 20. u mjesecu. "
              "Kredit se ne odobrava s počekom otplate."),
    ],
)

STAMBENI = (
    "Stambeni krediti - informativni list",
    [
        ("p", "Stambeni kredit Banke Demo namijenjen je kupnji, izgradnji, dogradnji ili adaptaciji stambene nekretnine "
              "u Republici Hrvatskoj te refinanciranju postojećeg stambenog kredita u drugoj banci."),
        ("h2", "Osnovne karakteristike"),
        ("ul", [
            "Valuta kredita: isključivo EUR.",
            "Iznos kredita: od 20.000,00 EUR do 400.000,00 EUR.",
            "Rok otplate: od 5 do 30 godina (od 60 do 360 mjeseci).",
            "Kamatna stopa: fiksna za cijelo vrijeme otplate.",
            "Otplata: u jednakim mjesečnim anuitetima.",
            "Banka financira do 85 % procijenjene vrijednosti nekretnine; ostatak je vlastito učešće klijenta.",
        ]),
        ("h2", "Tko može ugovoriti stambeni kredit"),
        ("p", "Stambeni kredit mogu ugovoriti isključivo državljani Republike Hrvatske. "
              "Osobe koje nisu državljani Republike Hrvatske ne mogu ugovoriti stambeni kredit u Banci Demo, "
              "ali mogu ugovoriti gotovinski nenamjenski kredit do 40.000,00 EUR."),
        ("ul", [
            "Klijent mora imati najmanje 18 godina, a kredit mora biti otplaćen do navršene 75. godine života.",
            "Potrebno je stalno zaposlenje ili redovita mirovina.",
            "Kamatna stopa ne ovisi o tome prima li klijent plaću u Banci Demo.",
        ]),
        ("h2", "Kamatne stope"),
        ("p", "Fiksne kamatne stope za stambene kredite kreću se od 3,49 % do 4,09 % godišnje, ovisno o roku otplate i iznosu kredita. "
              "Za veće iznose kredita i kraće rokove otplate kamatna stopa je niža."),
        ("h2", "Instrumenti osiguranja"),
        ("ul", [
            "Založno pravo (hipoteka) na nekretnini u korist Banke.",
            "Polica osiguranja nekretnine od požara i drugih opasnosti, vinkulirana u korist Banke.",
            "Zadužnica korisnika kredita.",
            "Polica životnog osiguranja nije obvezna.",
        ]),
        ("h2", "Potrebna dokumentacija"),
        ("ul", [
            "Zahtjev za stambeni kredit.",
            "Osobna iskaznica.",
            "Potvrda poslodavca o zaposlenju i visini primanja, ne starija od 30 dana.",
            "Kupoprodajni ugovor ili predugovor, odnosno građevinska dozvola i troškovnik za izgradnju ili adaptaciju.",
            "Izvadak iz zemljišne knjige, ne stariji od 30 dana.",
            "Procjena vrijednosti nekretnine od ovlaštenog procjenitelja s popisa Banke.",
        ]),
        ("h2", "Naknade i troškovi"),
        ("ul", [
            "Naknada za obradu kredita: bez naknade.",
            "Procjena nekretnine: plaća klijent ovlaštenom procjenitelju, okvirno od 200,00 do 400,00 EUR.",
            "Prijevremena otplata: bez naknade.",
            "Javnobilježnički troškovi i troškovi upisa hipoteke plaća klijent.",
        ]),
        ("h2", "Rok obrade zahtjeva"),
        ("p", "Odluka o zahtjevu za stambeni kredit donosi se u roku od 5 radnih dana od predaje potpune dokumentacije."),
    ],
)

TARIFA = (
    "Tarifa naknada za građane",
    [
        ("p", "Izvadak iz Odluke o naknadama za usluge Banke Demo d.d. za fizičke osobe. Svi iznosi su u EUR i vrijede od 1. rujna 2026."),
        ("h2", "Tekući račun"),
        ("table", [
            ["Usluga", "Naknada"],
            ["Otvaranje tekućeg računa", "bez naknade"],
            ["Mjesečno vođenje računa - paket Osnovni", "2,00 EUR"],
            ["Mjesečno vođenje računa - paket Plus (uključena debitna kartica i neograničene interne uplate)", "4,50 EUR"],
            ["Mjesečno vođenje računa za mlade do 27 godina i umirovljenike", "bez naknade"],
            ["Zatvaranje računa", "bez naknade"],
        ]),
        ("h2", "Platni promet"),
        ("table", [
            ["Usluga", "Internet i mobilno bankarstvo", "Poslovnica"],
            ["Nacionalna i SEPA uplata u korist računa u Banci Demo", "bez naknade", "1,00 EUR"],
            ["Nacionalna i SEPA uplata u korist računa u drugoj banci", "0,20 EUR", "2,00 EUR"],
            ["Trenutno plaćanje (instant)", "0,20 EUR", "2,00 EUR"],
            ["Plaćanje u stranoj valuti izvan SEPA područja", "0,3 %, min. 10,00 EUR", "0,5 %, min. 15,00 EUR"],
        ]),
        ("h2", "Kartice"),
        ("table", [
            ["Usluga", "Naknada"],
            ["Debitna kartica - godišnja članarina", "bez naknade uz paket Plus, inače 10,00 EUR"],
            ["Kreditna kartica Demo Classic - godišnja članarina", "25,00 EUR"],
            ["Podizanje gotovine na bankomatima Banke Demo", "bez naknade"],
            ["Podizanje gotovine na bankomatima drugih banaka u RH", "1,50 EUR"],
            ["Podizanje gotovine u inozemstvu", "2 %, min. 3,00 EUR"],
            ["Izrada nove kartice zbog gubitka ili krađe", "7,00 EUR"],
        ]),
        ("h2", "Ostale usluge"),
        ("table", [
            ["Usluga", "Naknada"],
            ["Potvrda o stanju računa", "5,00 EUR"],
            ["Opomena zbog neplaćenog dospjelog duga", "5,00 EUR"],
            ["Najam sefa - godišnje (mali)", "60,00 EUR"],
        ]),
    ],
)

FAQ = (
    "Česta pitanja",
    [
        ("h2", "Kako mogu blokirati izgubljenu ili ukradenu karticu?"),
        ("p", "Karticu odmah blokirajte u mobilnoj aplikaciji Demo mBank (Kartice > Blokiraj karticu) ili pozivom na broj "
              "0800 1234 koji je dostupan 0-24, svaki dan. Iz inozemstva nazovite +385 1 4000 123. Nova kartica izrađuje se u roku od 5 radnih dana."),
        ("h2", "Kako se prijaviti u mobilno bankarstvo?"),
        ("p", "Aplikaciju Demo mBank preuzmite s Google Play ili App Store trgovine. Za aktivaciju vam je potreban aktivacijski kod "
              "koji dobivate u poslovnici ili putem internet bankarstva. Aplikacija je besplatna za sve klijente s tekućim računom."),
        ("h2", "Mogu li otvoriti račun ako nisam državljanin Republike Hrvatske?"),
        ("p", "Da. Osobe koje nisu državljani Republike Hrvatske mogu otvoriti tekući račun uz važeću putovnicu ili osobnu iskaznicu "
              "i OIB. Za odobravanje kredita potrebna je i dozvola boravka u RH. Stambeni krediti dostupni su isključivo državljanima RH."),
        ("h2", "Koliko traje odobravanje gotovinskog kredita?"),
        ("p", "Ako su priložene sve potrebne isprave, odluka o gotovinskom kreditu donosi se isti ili idući radni dan. "
              "Postojeći klijenti koji primaju plaću u Banci Demo mogu zatražiti kredit do 15.000,00 EUR izravno u mobilnoj aplikaciji."),
        ("h2", "Mogu li prijevremeno otplatiti kredit?"),
        ("p", "Da, kredit možete djelomično ili u cijelosti otplatiti prije roka. Za stambene kredite prijevremena otplata je bez naknade. "
              "Za gotovinske kredite naknada je opisana u Općim uvjetima za gotovinske kredite."),
        ("h2", "Nudi li banka kredite u stranoj valuti?"),
        ("p", "Ne. Banka Demo odobrava kredite isključivo u eurima (EUR), bez valutne klauzule."),
        ("h2", "Što je kamatna stopa i je li fiksna?"),
        ("p", "Sve kamatne stope na kredite Banke Demo su fiksne, što znači da se mjesečna rata ne mijenja tijekom cijelog razdoblja otplate."),
        ("h2", "Kako mogu promijeniti PIN kartice?"),
        ("p", "PIN možete promijeniti na bilo kojem bankomatu Banke Demo, u izborniku Promjena PIN-a, bez naknade."),
    ],
)



# ---------------------------------------------------------------------------


FONT_CANDIDATES = [Path(r"C:\Windows\Fonts\arial.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
                   Path("/Library/Fonts/Arial Unicode.ttf")]


def _font():
    for p in FONT_CANDIDATES:
        if p.exists():
            return p
    raise SystemExit("Nije pronađen TTF font s hrvatskim slovima (npr. Arial).")


def write_pdf(doc, path):
    title, blocks = doc
    pdf = FPDF()
    pdf.add_font("Body", "", str(_font()))
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()
    pdf.set_font("Body", size=16)
    pdf.multi_cell(0, 9, f"Banka Demo d.d. - {title}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    for kind, content in blocks:
        if kind == "h2":
            pdf.ln(2)
            pdf.set_font("Body", size=13)
            pdf.multi_cell(0, 8, content, new_x="LMARGIN", new_y="NEXT")
        elif kind == "p":
            pdf.set_font("Body", size=10)
            pdf.multi_cell(0, 5.5, content, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        elif kind == "ul":
            pdf.set_font("Body", size=10)
            for item in content:
                pdf.multi_cell(0, 5.5, "- " + item, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
    pdf.output(str(path))


def write_docx(doc, path):
    title, blocks = doc
    d = Document()
    d.add_heading(title, level=1)
    for kind, content in blocks:
        if kind == "h2":
            d.add_heading(content, level=2)
        elif kind == "p":
            d.add_paragraph(content)
        elif kind == "ul":
            for item in content:
                d.add_paragraph(item, style="List Bullet")
        elif kind == "table":
            table = d.add_table(rows=0, cols=len(content[0]))
            table.style = "Table Grid"
            for row in content:
                cells = table.add_row().cells
                for c, value in zip(cells, row):
                    c.text = value
    d.save(str(path))


def main():
    for sub in ("pdf", "docx"):
        (RAW_DIR / sub).mkdir(parents=True, exist_ok=True)
    write_pdf(GOTOVINSKI, RAW_DIR / "pdf" / "opci-uvjeti-gotovinski-krediti.pdf")
    write_pdf(STAMBENI, RAW_DIR / "pdf" / "stambeni-krediti-informativni-list.pdf")
    write_docx(TARIFA, RAW_DIR / "docx" / "tarifa-naknada-gradani.docx")
    write_docx(FAQ, RAW_DIR / "docx" / "cesta-pitanja.docx")
    print("Generirano u", RAW_DIR)


if __name__ == "__main__":
    main()
